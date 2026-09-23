import io
from datetime import datetime, date
from typing import List, Dict, Any, Tuple
import openpyxl
from app.imports.parsers.base import BaseStatementParser
from app.imports.exceptions import ParseError


class XLSXStatementParser(BaseStatementParser):
    """
    Excel (XLSX) statement parser using openpyxl in read_only mode.
    Selects the primary data sheet, extracts headers, and normalizes cell values.
    """

    def parse_bytes(
        self, file_bytes: bytes, filename: str
    ) -> Tuple[List[str], List[Dict[str, Any]]]:
        if not file_bytes:
            raise ParseError(f"Uploaded Excel file '{filename}' is empty.")

        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(file_bytes), read_only=True, data_only=True
            )
        except Exception as e:
            raise ParseError(f"Failed to open Excel workbook '{filename}': {str(e)}")

        sheetnames = wb.sheetnames
        if not sheetnames:
            raise ParseError(f"Excel workbook '{filename}' has no worksheets.")

        # Find first non-empty sheet
        active_sheet = None
        for name in sheetnames:
            ws = wb[name]
            # Peek up to 5 rows
            has_data = False
            for row in ws.iter_rows(max_row=5, values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    has_data = True
                    break
            if has_data:
                active_sheet = ws
                break

        if active_sheet is None:
            active_sheet = wb[sheetnames[0]]

        # Read rows
        raw_rows = list(active_sheet.iter_rows(values_only=True))
        if not raw_rows:
            raise ParseError(f"Worksheet '{active_sheet.title}' in '{filename}' is completely empty.")

        # Find header row
        header_idx = -1
        headers: List[str] = []
        for idx, row in enumerate(raw_rows):
            cleaned = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cleaned:
                header_idx = idx
                headers = [str(c).strip() if c is not None else "" for c in row]
                # Trim trailing empty header cells
                while headers and not headers[-1]:
                    headers.pop()
                break

        if header_idx == -1 or not headers:
            raise ParseError(f"Could not locate a header row in Excel file '{filename}'.")

        # Sanitize headers
        sanitized_headers = []
        seen = {}
        for col_i, h in enumerate(headers):
            clean_h = h.strip()
            if not clean_h:
                clean_h = f"Column_{col_i + 1}"
            if clean_h in seen:
                seen[clean_h] += 1
                sanitized_headers.append(f"{clean_h}_{seen[clean_h]}")
            else:
                seen[clean_h] = 1
                sanitized_headers.append(clean_h)

        data_rows: List[Dict[str, Any]] = []
        for line_num, row in enumerate(raw_rows[header_idx + 1 :], start=header_idx + 2):
            if not row or not any(v is not None and str(v).strip() != "" for v in row):
                continue  # skip empty

            row_dict: Dict[str, Any] = {"__row_number__": line_num}
            for col_idx, col_name in enumerate(sanitized_headers):
                cell_val = row[col_idx] if col_idx < len(row) else None
                # Format datetime/date cells nicely
                if isinstance(cell_val, (datetime, date)):
                    cell_val = cell_val.strftime("%Y-%m-%d")
                elif cell_val is not None:
                    cell_val = str(cell_val).strip()
                else:
                    cell_val = ""
                row_dict[col_name] = cell_val

            data_rows.append(row_dict)

        if not data_rows:
            raise ParseError(f"Excel file '{filename}' has headers but no transaction rows.")

        return sanitized_headers, data_rows
